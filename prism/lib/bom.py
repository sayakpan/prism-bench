'''
BOM (Bill of Materials) Excel parsing for Moodboard Style.

A Linctex BOM export carries a "BOM - Fabrics" table (header row with columns
Name / Roll Width / Height / Avg. Fabric Length / Piece / Efficiency ...). We read
the first fabric row of the first sheet that has this table and derive:

    consumption      = (roll_width * avg_fabric_length * gsm) / divisor
    marker_efficiency = the fabric row's Efficiency

The divisor converts the fabric area to m^2 so that, multiplied by GSM (g/m^2), it
yields grams — and it depends on the unit the length cells are expressed in:

    cm    -> 10000   (1 m^2 = 10000 cm^2)   e.g. "155.00cm"
    inch  -> 1550    (1 m^2 = 1550 in^2)    e.g. "61.024in"

The unit is read from the length cells' own suffix. If no unit can be determined
(bare numbers, no cm/in suffix), consumption is left blank rather than guessed, so
a wrong divisor is never silently applied.

GSM is parsed out of the fabric Name cell (e.g. "US-1514 GSM-165 50_ ORGANIC
COTTON ..." -> 165), so parsing is self-contained and works no matter how the file
was uploaded. Numbers tolerate unit suffixes ("155.00cm" -> 155.0).
'''

import io
import re

import openpyxl

# Column-label matchers for the "BOM - Fabrics" header row. A header cell matches
# a canonical key when its normalised text contains the given token(s).
_FABRIC_COLS = {
    'name': ('name',),
    'roll_width': ('roll', 'width'),
    'avg_length': ('avg', 'length'),   # "Avg. Fabric Length / Piece"
    'efficiency': ('efficiency',),
}

# Area-to-m^2 divisors keyed by the length unit the BOM uses. GSM is g/m^2, so
# (roll * length) must be converted to m^2 before multiplying by GSM. A unit we
# can't identify maps to no divisor -> consumption is left blank.
_UNIT_DIVISORS = {
    'cm': 10000,   # 1 m^2 = 10000 cm^2
    'in': 1550,    # 1 m^2 = 1550 in^2
}

# Rows to scan (per sheet) when hunting for the fabrics header — the table sits
# near the top; no need to walk huge sheets.
_HEADER_SCAN_ROWS = 40

# Rows below the header to search for the first fabric data row.
_DATA_SCAN_ROWS = 50


def gsm_from_text(text):
    ''' First GSM value in a free-text string ("... GSM-165 ...", "165 GSM"), or None. '''
    if not text:
        return None
    s = str(text)
    m = re.search(r'GSM[\s\-:_]*(\d+(?:\.\d+)?)', s, re.IGNORECASE)
    if not m:
        m = re.search(r'(\d+(?:\.\d+)?)\s*GSM', s, re.IGNORECASE)
    return float(m.group(1)) if m else None


def parse_bom_metrics(content):
    '''
    Parse consumption + marker efficiency from BOM Excel bytes. Returns a dict
    (values may be None where a field is missing) or None if the file has no
    recognisable fabrics table. Never raises for content issues — callers treat a
    None/partial result as "nothing to auto-fill".

    Returns: {
        'consumption', 'marker_efficiency', 'gsm',
        'roll_width', 'avg_length', 'fabric_name',
    }
    '''
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            header = _find_fabrics_header(ws)
            if not header:
                continue
            row_idx, cols = header
            data = _first_data_row(ws, row_idx, cols)
            if not data:
                continue

            roll = _num(data.get('roll_width'))
            avg = _num(data.get('avg_length'))
            eff = _num(data.get('efficiency'))
            name = data.get('name')
            gsm = gsm_from_text(name)

            # Unit comes from the length cells' suffix; the two should agree, so
            # take roll width's and fall back to avg length's.
            unit = _length_unit(data.get('roll_width')) or _length_unit(data.get('avg_length'))
            divisor = _UNIT_DIVISORS.get(unit)

            consumption = None
            if roll and avg and gsm and divisor:
                consumption = round((roll * avg * gsm) / divisor, 3)

            return {
                'consumption': consumption,
                'marker_efficiency': eff,
                'gsm': gsm,
                'roll_width': roll,
                'avg_length': avg,
                'fabric_name': name,
            }
        return None
    finally:
        wb.close()


# --- helpers ---

def _norm(value):
    return re.sub(r'\s+', ' ', str(value).strip().lower()) if value not in (None, '') else ''


def _num(value):
    ''' Leading number in a cell, ignoring unit suffixes/thousands separators. '''
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = re.search(r'[-+]?\d*\.?\d+', str(value).replace(',', ''))
    return float(m.group()) if m else None


def _length_unit(value):
    '''
    Length unit implied by a cell's suffix: 'in' for inch ("61.024in", '45.7"'),
    'cm' for centimetre ("155.00cm"), or None when no unit is present (e.g. a bare
    number, whose divisor we won't guess). The suffix directly follows the number,
    so we anchor on a preceding digit; inch is checked first.
    '''
    if value is None:
        return None
    s = str(value).lower()
    if re.search(r'\d\s*(?:inch|in|["″])', s):
        return 'in'
    if re.search(r'\d\s*cm', s):
        return 'cm'
    return None


def _find_fabrics_header(ws):
    '''
    Locate the fabrics table header on a sheet: the first row (within the top
    _HEADER_SCAN_ROWS) whose cells include both a Roll Width and an Efficiency
    label. Returns (row_index, {canonical_key: column_index}) or None.
    '''
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS), start=1):
        mapping = {}
        for cell in row:
            label = _norm(cell.value)
            if not label:
                continue
            for key, tokens in _FABRIC_COLS.items():
                if key not in mapping and all(tok in label for tok in tokens):
                    mapping[key] = cell.column  # 1-based column index
        if 'roll_width' in mapping and 'efficiency' in mapping:
            return r_idx, mapping
    return None


def _first_data_row(ws, header_row, cols):
    '''
    First populated data row beneath the header — the row that carries a Name (or,
    failing that, any mapped column). Reads via iter_rows (not ws.max_row, which
    is unreliable in read-only mode) and maps by real column index so it tolerates
    sheets that don't start at column A. Returns {canonical_key: value}.
    '''
    name_col = cols.get('name')
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + _DATA_SCAN_ROWS):
        by_col = {cell.column: cell.value for cell in row if cell.value not in (None, '')}
        if not by_col:
            continue
        values = {key: by_col.get(idx) for key, idx in cols.items()}
        anchor = values.get('name') if name_col else None
        if (anchor not in (None, '')) or any(v not in (None, '') for v in values.values()):
            return values
    return None
